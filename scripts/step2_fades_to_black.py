#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""第 2 步：把 ``Fades to black(ms)`` 展開成獨立的黑燈 cue。

燈流表上的 ``Fades to black(ms)`` 是「這個 cue 結束後花多久暗下來」，
但 QLC+ 的 Sequence 沒有這個概念——它只有每個 Step 自己的
FadeIn / Hold / FadeOut。所以這一步做兩件事：

1. 欄位名稱 ``Fades to black(ms)`` 改成 ``Fade out(ms)``（QLC+ 認得的名字）。
2. 只要某一列的 ``Fades to black(ms)`` 有值（非空且非 0），
   就在它下面補一個**全通道歸零的黑燈 cue**，
   該 cue 的 ``Fade in(ms)`` 等於上一列的 ``Fade out(ms)``，
   ``Hold(ms)`` 與 ``Fade out(ms)`` 都是 0。
   暗場就發生在這段淡入裡：畫面在這段時間內從上一個 cue 漸暗到全黑。

產物會放在原檔旁邊：

    1_raw.xlsx  →  2_forqxw.xlsx

用法::

    python3 step2_fades_to_black.py temp_AllCueList        # 整個專案資料夾
    python3 step2_fades_to_black.py 某某/1_raw.xlsx          # 單一檔案
    python3 step2_fades_to_black.py temp_AllCueList --dry-run
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional

try:
    import openpyxl
except ImportError:                                     # pragma: no cover
    print("需要 openpyxl：python3 -m pip install openpyxl", file=sys.stderr)
    raise SystemExit(2)


SUFFIX = "_forqxw"
RAW_STEM = "1_raw"                  # 第 1 步拆出來的檔名
FORQXW_NAME = "2_forqxw.xlsx"       # 這一步的產物
HEADER_ROW = 2          # 第 1 列是標題（例如「60RC」），第 2 列才是欄位名稱

FADE_TO_BLACK = "fadestoblackms"
FADE_OUT_HEADER = "Fade out(ms)"

# 欄位名稱 -> 角色。其餘有名稱的欄位一律視為燈具通道欄。
META_HEADERS = {
    "#": "index", "no": "index", "step": "index", "序號": "index",
    "fadein": "fade_in", "fadeinms": "fade_in", "淡入": "fade_in",
    "hold": "hold", "holdms": "hold", "duration": "hold", "durationms": "hold",
    "時間": "hold",
    "fadeout": "fade_out", "fadeoutms": "fade_out", "淡出": "fade_out",
    FADE_TO_BLACK: "fade_to_black", "fadetoblackms": "fade_to_black",
    "fadestoblack": "fade_to_black", "fadetoblack": "fade_to_black",
    "note": "note", "notes": "note", "備註": "note",
}


def normalize(text) -> str:
    """去掉空白/標點並轉小寫，用於寬鬆比對欄位名稱。"""
    return re.sub(r"[\s\-_()（）./﻿]+", "", str(text or "").strip().lower())


def to_number(value) -> Optional[float]:
    """把儲存格內容轉成數字，轉不動就回 None。"""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def classify_columns(header: List) -> Dict[str, object]:
    """把表頭分成各種角色欄與燈具通道欄（都用 0-based 欄號）。"""
    roles: Dict[str, int] = {}
    channels: List[int] = []
    for index, cell in enumerate(header):
        title = str(cell).strip() if cell is not None else ""
        if not title:
            continue
        role = META_HEADERS.get(normalize(title))
        if role:
            roles.setdefault(role, index)
        else:
            channels.append(index)
    return {"roles": roles, "channels": channels}


def transform_sheet(source, target) -> int:
    """把一張工作表複製到 target，並展開黑燈 cue。回傳插入了幾列。"""
    rows = [list(row) for row in source.iter_rows(values_only=True)]
    if len(rows) < HEADER_ROW:
        for row in rows:
            target.append(row)
        return 0

    header = list(rows[HEADER_ROW - 1])
    layout = classify_columns(header)
    roles: Dict[str, int] = layout["roles"]
    channels: List[int] = layout["channels"]

    black_col = roles.get("fade_to_black")
    index_col = roles.get("index")
    fade_in_col = roles.get("fade_in")
    hold_col = roles.get("hold")

    # 欄位改名：Fades to black(ms) -> Fade out(ms)
    if black_col is not None:
        header[black_col] = FADE_OUT_HEADER

    for row in rows[:HEADER_ROW - 1]:
        target.append(row)
    target.append(header)

    width = max((len(row) for row in rows), default=len(header))
    width = max(width, len(header))
    inserted = 0
    number = 0

    for row in rows[HEADER_ROW:]:
        row = list(row) + [None] * (width - len(row))
        number += 1
        if index_col is not None and index_col < len(row):
            row[index_col] = number
        target.append(row)

        if black_col is None or black_col >= len(row):
            continue
        fade = to_number(row[black_col])
        if not fade:                    # 空白或 0 都不用補黑燈 cue
            continue

        # 補一列全通道歸零的黑燈 cue：Fade in 等於上一列的 Fade out
        # （暗場就發生在這段淡入裡），Hold 與 Fade out 都是 0
        blackout: List = [None] * width
        number += 1
        if index_col is not None:
            blackout[index_col] = number
        if fade_in_col is not None:
            blackout[fade_in_col] = row[black_col]
        if hold_col is not None:
            blackout[hold_col] = 0
        blackout[black_col] = 0
        for column in channels:
            blackout[column] = 0
        target.append(blackout)
        inserted += 1

    return inserted


def save_hint(out_path: Path) -> str:
    """依平台給出對應的排除建議。"""
    if sys.platform == "darwin":
        return (f"    舊檔多半是用終端機或別的程式產生的，macOS 不讓未簽章的 App 改它。\n"
                f"    請把 {out_path.parent.parent} 整個資料夾刪掉後重跑，\n"
                f"    或到「系統設定 → 隱私權與安全性 → 完全取用磁碟」把這個 App 加進去。")
    if sys.platform.startswith("win"):
        return (f"    最常見的原因是這個檔案正開在 Excel 裡。請把它關掉後重跑，\n"
                f"    或把 {out_path.parent.parent} 整個資料夾刪掉。")
    return f"    請確認檔案沒有被其他程式開著，或把 {out_path.parent.parent} 刪掉後重跑。"


class SaveBlocked(Exception):
    """存檔被作業系統擋下，且刪不掉舊檔。"""


def save_workbook(book, out_path: Path) -> None:
    """存檔，並處理 macOS 擋下「覆寫別的程式建立的舊檔」的情況。

    打包出來的 .app 是未簽章程式，在 ~/Documents、~/Desktop 這類受 TCC
    保護的位置，覆寫「由終端機或其他程式產生」的舊檔會得到 EPERM
    （Operation not permitted），但在同一個資料夾裡「建立新檔」是允許的。
    所以先刪掉舊檔再寫；真的連刪都刪不掉才報錯。
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        book.save(out_path)
        return
    except PermissionError:
        pass

    try:
        out_path.unlink()
    except FileNotFoundError:
        pass
    except OSError as exc:
        raise SaveBlocked(
            f"沒有權限覆寫 {out_path}（{exc.strerror}）。\n" + save_hint(out_path)
        ) from exc

    book.save(out_path)


def convert_workbook(path: Path, *, dry_run: bool = False) -> Optional[Path]:
    """把一份 cuelist 活頁簿轉成 ``…_forqxw.xlsx``。"""
    out_path = path.with_name(FORQXW_NAME if path.stem == RAW_STEM
                              else f"{path.stem}{SUFFIX}{path.suffix}")
    source = openpyxl.load_workbook(path, data_only=True)
    target = openpyxl.Workbook()
    target.remove(target.active)

    total = 0
    for sheet in source.worksheets:
        new_sheet = target.create_sheet(title=sheet.title)
        total += transform_sheet(sheet, new_sheet)
        for letter, dimension in sheet.column_dimensions.items():
            if dimension.width:
                new_sheet.column_dimensions[letter].width = dimension.width

    note = f"（補了 {total} 個黑燈 cue）" if total else "（沒有要補的黑燈 cue）"
    # 檔名每個資料夾都一樣，所以訊息帶上所屬資料夾才分得出是哪一組
    where = f"{path.parent.name}/" if path.parent.name else ""
    prefix = "  [試跑] " if dry_run else "  "
    print(f"{prefix}{where}{path.name} → {out_path.name} {note}")
    if not dry_run:
        save_workbook(target, out_path)
    return out_path


def find_workbooks(root: Path) -> List[Path]:
    """收集要轉換的 cuelist 檔（排除已經轉好的與 Excel 暫存鎖定檔）。"""
    if root.is_file():
        return [root]
    found = [
        path for path in sorted(root.rglob("*.xlsx"))
        if not path.name.startswith("~$") and not path.stem.endswith(SUFFIX)
    ]
    return found


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="把 Fades to black(ms) 展開成獨立的黑燈 cue，"
                    "另存成 …_forqxw.xlsx",
        formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", type=Path,
                        help="專案資料夾（會遞迴找 .xlsx）或單一 .xlsx")
    parser.add_argument("--dry-run", action="store_true",
                        help="只顯示會做什麼，不實際寫檔")
    args = parser.parse_args(argv)

    if not args.source.exists():
        print(f"找不到來源：{args.source}", file=sys.stderr)
        return 2

    workbooks = find_workbooks(args.source)
    if not workbooks:
        print(f"{args.source} 底下沒有可以轉換的 .xlsx", file=sys.stderr)
        return 1

    for path in workbooks:
        try:
            convert_workbook(path, dry_run=args.dry_run)
        except SaveBlocked as exc:
            print(f"  [失敗] {exc}", file=sys.stderr)
            return 1
        except Exception as exc:                        # noqa: BLE001
            print(f"  ⚠ 轉換失敗 {path.name}: {exc}", file=sys.stderr)
            return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
