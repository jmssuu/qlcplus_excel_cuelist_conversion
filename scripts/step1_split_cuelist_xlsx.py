#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把一份總表 .xlsx 拆成每張工作表各自的 cuelist .xlsx。

用法::

    python3 step1_split_cuelist_xlsx.py Project1.xlsx
    python3 step1_split_cuelist_xlsx.py Project1.xlsx --start-sheet 4 --outdir .
    python3 step1_split_cuelist_xlsx.py Project1.xlsx /path/to/Music Basic_stage.qxw

產出結構（以 Project1.xlsx 的第 5 張工作表為例）::

    temp_Project1/XX組-表演名稱/1_raw.xlsx

轉換規則：
* 從第 ``--start-sheet`` 張工作表（預設第 4 張）開始，每張工作表各自產生一個資料夾。
* 在工作表裡找每一個內容為 ``#`` 的儲存格，以它為左上角往外框出一個區塊：
  往上一列（區塊標題列）、往下數到編號中斷（空白）為止、往右數到 ``Note`` 欄為止。
* 每個 ``#`` 區塊在輸出檔裡各成一張工作表，名稱為「原工作表名稱 + ``#`` 上方那格的文字」。
* 第二個參數是音樂資料夾的路徑（省略時用執行目錄下的 ``Music``），會把裡面
  「與工作表同名的 .mp3」複製進該工作表的資料夾。
* 第三個參數是底稿 .qxw（可省略），會複製到產出的專案資料夾根目錄，
  供 ``step3_cuelist_to_qxw.py`` 當成 merge 的底稿。複製前會先刪掉該資料夾根目錄
  既有的 .qxw（不含子資料夾），免得上次留下的底稿被誤當成這次的底稿。
"""

from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path
from typing import List, Optional, Tuple

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("需要 openpyxl，請先執行：python3 -m pip install openpyxl")

# 拆出來的原始表。資料夾本身已經是工作表名稱，檔名不必再重複一次；
# 前面的數字對應流程步驟，排序就是流程順序。
RAW_NAME = "1_raw.xlsx"

MARKER = "#"
END_HEADER = "note"
INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_NAME = 31


def cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def is_number(value) -> bool:
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    text = cell_text(value)
    if not text:
        return False
    try:
        float(text)
    except ValueError:
        return False
    return True


def find_markers(grid: List[List]) -> List[Tuple[int, int]]:
    """回傳所有 ``#`` 儲存格的 (列, 欄) 索引（0 起算），由上到下、由左到右。"""
    found = []
    for r, row in enumerate(grid):
        for c, value in enumerate(row):
            if cell_text(value) == MARKER:
                found.append((r, c))
    return found


def block_bounds(grid: List[List], row: int, col: int) -> Tuple[int, int, int, int]:
    """以 ``#`` 的位置算出區塊範圍 (起始列, 結束列, 起始欄, 結束欄)，皆含端點。"""
    top = row - 1 if row > 0 else row

    bottom = row
    for r in range(row + 1, len(grid)):
        line = grid[r]
        value = line[col] if col < len(line) else None
        if not is_number(value):
            break
        bottom = r

    header = grid[row]
    right = col
    for c in range(col + 1, len(header)):
        right = c
        if cell_text(header[c]).lower() == END_HEADER:
            break
    return top, bottom, col, right


def block_title(grid: List[List], row: int, col: int) -> str:
    """``#`` 上方那格的文字；整列往右找第一個非空白，避免標題被合併儲存格擠開。"""
    if row == 0:
        return ""
    above = grid[row - 1]
    if col < len(above) and cell_text(above[col]):
        return cell_text(above[col])
    for value in above[col:]:
        if cell_text(value):
            return cell_text(value)
    return ""


def unique_sheet_name(base: str, used: set) -> str:
    name = INVALID_SHEET_CHARS.sub("-", base).strip() or "Sheet"
    name = name[:MAX_SHEET_NAME]
    if name not in used:
        used.add(name)
        return name
    for n in range(2, 1000):
        suffix = f"_{n}"
        candidate = name[: MAX_SHEET_NAME - len(suffix)] + suffix
        if candidate not in used:
            used.add(candidate)
            return candidate
    raise ValueError(f"無法為 {base} 產生唯一的工作表名稱")


def resolve_music_dir(path: Path) -> Optional[Path]:
    """確認音樂資料夾存在；同層若只有大小寫不同（Music/music）也接受。"""
    if path.is_dir():
        return path
    parent = path.parent if str(path.parent) else Path(".")
    if parent.is_dir():
        for entry in sorted(parent.iterdir()):
            if entry.is_dir() and entry.name.lower() == path.name.lower():
                return entry
    return None


def find_music_file(music_dir: Path, sheet_title: str) -> Optional[Path]:
    """在音樂資料夾裡找與工作表同名的 .mp3；找不到再試不分大小寫、去掉頭尾空白。"""
    wanted = sheet_title.strip()
    exact = music_dir / f"{wanted}.mp3"
    if exact.is_file():
        return exact
    for entry in sorted(music_dir.iterdir()):
        if entry.is_file() and entry.suffix.lower() == ".mp3" and entry.stem.strip().lower() == wanted.lower():
            return entry
    return None


def safe_dir_name(name: str) -> str:
    cleaned = name.replace("/", "-").replace("\\", "-").strip()
    return cleaned or "sheet"


def convert_sheet(worksheet, out_path: Path) -> int:
    """把一張工作表裡的所有 ``#`` 區塊寫成一個新的 .xlsx，回傳區塊數量。"""
    grid = [list(row) for row in worksheet.iter_rows(values_only=True)]
    markers = find_markers(grid)
    if not markers:
        return 0

    book = openpyxl.Workbook()
    book.remove(book.active)
    used: set = set()
    for row, col in markers:
        top, bottom, left, right = block_bounds(grid, row, col)
        title = block_title(grid, row, col)
        sheet = book.create_sheet(unique_sheet_name(worksheet.title + title, used))
        for r in range(top, bottom + 1):
            line = grid[r]
            sheet.append([line[c] if c < len(line) else None for c in range(left, right + 1)])

    save_workbook(book, out_path)
    return len(markers)


def save_hint(out_path: Path) -> str:
    """依平台給出對應的排除建議。"""
    if sys.platform == "darwin":
        return (f"    舊檔多半是用終端機或別的程式產生的，macOS 不讓未簽章的 App 改它。\n"
                f"    請把 {out_path.parent.parent} 整個資料夾刪掉後重跑，\n"
                f"    或到「系統設定 → 隱私權與安全性 → 完全取用磁碟」把這個 App 加進去。")
    if sys.platform.startswith("win"):
        return (f"    最常見的原因是這個檔案正開在 Excel 裡。請把它關掉後重跑，\n"
                f"    或把 {out_path.parent.parent} 整個資料夾刪掉。")
    return f"    請確認檔案沒有被其他程式開著，或把 {out_path.parent.parent} 刪掉後重跑。"


class SaveBlocked(Exception):
    """存檔被作業系統擋下，且刪不掉舊檔。"""


def save_workbook(book, out_path: Path) -> None:
    """存檔，並處理 macOS 擋下「覆寫別的程式建立的舊檔」的情況。

    打包出來的 .app 是未簽章程式，在 ~/Documents、~/Desktop 這類受 TCC
    保護的位置，覆寫「由終端機或其他程式產生」的舊檔會得到 EPERM
    （Operation not permitted），但在同一個資料夾裡「建立新檔」是允許的。
    所以先刪掉舊檔再寫；真的連刪都刪不掉才報錯。
    """
    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        book.save(out_path)
        return
    except PermissionError:
        pass

    try:
        out_path.unlink()
    except FileNotFoundError:
        pass
    except OSError as exc:
        raise SaveBlocked(
            f"沒有權限覆寫 {out_path}（{exc.strerror}）。\n" + save_hint(out_path)
        ) from exc

    book.save(out_path)


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="把總表 .xlsx 拆成每張工作表的 cuelist .xlsx")
    parser.add_argument("source", type=Path, help="要轉換的 .xlsx")
    parser.add_argument("--start-sheet", type=int, default=4,
                        help="從第幾張工作表開始轉換（1 起算，預設 4）")
    parser.add_argument("--outdir", type=Path, default=Path("."),
                        help="輸出根目錄，預設為目前執行目錄")
    parser.add_argument("music_dir", type=Path, nargs="?", default=Path("Music"),
                        help="存放 mp3 的資料夾路徑，預設為執行目錄下的 Music")
    parser.add_argument("qxw", type=Path, nargs="?", default=None,
                        help="要複製進輸出專案資料夾的 .qxw 檔案路徑")
    args = parser.parse_args(argv)

    if not args.source.is_file():
        parser.error(f"找不到檔案：{args.source}")
    if args.start_sheet < 1:
        parser.error("--start-sheet 必須 >= 1")
    if args.qxw is not None:
        if not args.qxw.is_file():
            parser.error(f"找不到 .qxw 檔案：{args.qxw}")
        if args.qxw.suffix.lower() != ".qxw":
            parser.error(f"底稿必須是 .qxw：{args.qxw}")

    book = openpyxl.load_workbook(args.source, data_only=True, read_only=True)
    try:
        sheets = book.worksheets[args.start_sheet - 1:]
        if not sheets:
            print(f"{args.source} 只有 {len(book.worksheets)} 張工作表，第 "
                  f"{args.start_sheet} 張之後沒有東西可以轉換。")
            return 1

        music_dir = resolve_music_dir(args.music_dir)
        if music_dir is None:
            print(f"[提醒] 找不到音樂資料夾 {args.music_dir}，略過複製 mp3。")

        root = args.outdir / f"temp_{args.source.stem}"
        if args.qxw is not None:
            root.mkdir(parents=True, exist_ok=True)
            target = root / args.qxw.name
            source_qxw = args.qxw.resolve()
            for stale in sorted(root.glob("*.qxw")):
                if stale.is_file() and stale.resolve() != source_qxw:
                    stale.unlink()
                    print(f"[清除] 移除舊的 {stale}")
            if source_qxw != target.resolve():
                shutil.copy2(args.qxw, target)
            print(f"[OK] 複製底稿 {args.qxw.name} -> {root}")

        for worksheet in sheets:
            folder = root / safe_dir_name(worksheet.title)
            out_path = folder / RAW_NAME
            try:
                count = convert_sheet(worksheet, out_path)
            except SaveBlocked as exc:
                print(f"[失敗] {exc}", file=sys.stderr)
                return 1
            if count:
                print(f"[OK] {worksheet.title}：{count} 個區塊 -> {out_path}")
            else:
                print(f"[跳過] {worksheet.title}：找不到 '{MARKER}' 區塊")

            if music_dir is not None:
                mp3 = find_music_file(music_dir, worksheet.title)
                if mp3 is None:
                    print(f"[提醒] {music_dir}/ 裡找不到 {worksheet.title}.mp3")
                else:
                    folder.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(mp3, folder / mp3.name)
                    print(f"[OK] 複製音樂 {mp3.name} -> {folder}")
    finally:
        book.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
